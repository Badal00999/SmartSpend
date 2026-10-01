"""Generate the submission report. Optional tooling: pip install python-docx.
Run from any directory: python scripts/make_week5_report.py [output.docx]
"""
from pathlib import Path
import json, sys
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]
OUT = Path(sys.argv[1]) if len(sys.argv)>1 else ROOT.parent/'SpendWise_Week5_Report_Badal.docx'
doc=Document()
sec=doc.sections[0]
sec.top_margin=Inches(.62);sec.bottom_margin=Inches(.6)
sec.left_margin=sec.right_margin=Inches(.7)
normal=doc.styles['Normal'];normal.font.name='Calibri';normal.font.size=Pt(10)
normal.paragraph_format.space_after=Pt(6)
normal.paragraph_format.line_spacing=1.08
for style,size in [('Title',26),('Heading 1',17),('Heading 2',12)]:
 doc.styles[style].font.name='Calibri';doc.styles[style].font.size=Pt(size);doc.styles[style].font.color.rgb=RGBColor.from_string('0F766E')
head=sec.header.paragraphs[0];head.text='SPENDWISE  /  WEEK 5                                      TEST • DEBUG • MEASURE'
head.style='Caption'
foot=sec.footer.paragraphs[0];foot.text='Badal  |  YUVA Internship                                      Page '
fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');foot._p.append(fld)
def p(s,style=None): return doc.add_paragraph(s,style)
def h(s): doc.add_heading(s,1)
def sub(s): doc.add_heading(s,2)
def bullet(s): p(s,'List Bullet')
def code(s):
 r=p(s).runs[0];r.font.name='Consolas';r.font.size=Pt(8)
def table(headers,rows):
 t=doc.add_table(rows=1, cols=len(headers));t.style='Light Shading Accent 1'
 for c,v in zip(t.rows[0].cells,headers):c.text=str(v)
 for row in rows:
  for c,v in zip(t.add_row().cells,row):c.text=str(v)
 for row in t.rows:
  for c in row.cells:
   for para in c.paragraphs:
    para.paragraph_format.space_after=Pt(3)
    for run in para.runs:run.font.size=Pt(9)
 return t
def page():doc.add_page_break()

p('WEEK 5 / ENGINEERING REPORT','Subtitle')
doc.add_heading('SpendWise\nTesting, Debugging & Optimization',0)
p('Badal  •  Junior Full Stack Developer internship  •  26 September 2026')
p('Repository: https://github.com/Badal00999/SmartSpend\nDelivery folder: week5-testing-optimization/')
h('1. Scope and verified outcome')
p('The Week 4 React, Express and MongoDB application was copied into a separate Week 5 folder. The objective was to verify existing behaviour, expose missed boundary cases, fix observed defects and optimize a measured bottleneck. Earlier week folders were not modified. This report focuses on reproducible findings rather than repeating the application feature list.')
table(['Test layer','Baseline / inherited','Final','New'],[
 ['Server: Vitest + Supertest','76 freshly passed','100 passed','24'],
 ['Client: Vitest + Testing Library','78 freshly passed','92 passed','14'],
 ['Chromium: Playwright','28 inherited; not a pre-edit rerun','30 passed','2'],
 ['Total','182 existing cases','222 passed','40'],
])
p('Final client lint: 0 errors and 0 warnings. Production build: passed. Results are local executions, not a claim of GitHub CI success. Historical red tests, final logs, per-test JSON, source changes and raw query measurements are included in the ZIP. Extracting the ZIP into a fresh folder also passed locked setup, all 192 server/client tests and build; E2E was not repeated in that extracted copy.')
sub('Evidence-led workflow')
p('First, the unchanged server/client suites were executed. Next, new boundary tests were run against the old implementation and their failing output saved. The fixes were then applied and the same tests rerun, followed by the complete regression suites. A separate database experiment compared old and new indexes on an identical deterministic fixture.')
p('The initial new tests produced 14 server-validation failures, 2 transport failures, 3 form-validation failures and 4 logging failures. These are assertion counts, not 23 separate bugs. The next sections group them into four concrete defect cases and one performance finding.')

page();h('2. D1 — Calendar validation and financial integrity')
sub('Reproduction and observed behaviour')
p('POST /api/v1/transactions with date 2025-02-29 returned 201 instead of 422. PATCH with 2026-04-31 returned 200; a bulk payload containing an invalid date returned 201. JavaScript normalizes 2025-02-29 to 2025-03-01, silently changing the intended financial date.')
code("new Date('2025-02-29T00:00:00Z').toISOString()\n// '2025-03-01T00:00:00.000Z'")
p('For an invalid month filter (from=2026-13-01), /transactions and /stats/summary returned 400 downstream while /stats/by-category returned 200. The expected contract is consistent schema-level 422 responses before database access. The original schemas also accepted reversed date and amount ranges.')
sub('Root cause and exact fix')
p('server/src/validators/schemas.js checked a date regular expression and whether Date produced a non-NaN value. That does not reject calendar rollover. Query bounds checked only the expression. The front-end pure validator checked missing/future dates, but not calendar validity.')
code("const parsed = new Date(`${s}T00:00:00.000Z`)\nreturn Number.isFinite(parsed.getTime()) &&\n  parsed.toISOString().slice(0, 10) === s")
p('The server now requires exact UTC date round-tripping, validates ordered date/amount bounds, and applies the no-future rule only to transaction writes. Future query bounds remain allowed. TransactionForm.jsx mirrors calendar validation so client-side rules do not contradict the API.')
sub('Verification and rationale')
p('18 server cases cover impossible dates, the valid leap day, query bounds and write integrity. POST and bulk tests assert no insertion; PATCH asserts the original date is unchanged. Seven front-end cases check calendar boundaries, error-to-input ARIA linkage, focus, blocked submission and payload normalization. Native date inputs help ordinary typing but do not replace server validation.')
p('Two new browser tests verify that 29 February 2024 can be saved through the UI and survives reload, and that bypassing the input with same-origin browser fetch still rejects 29 February 2025 without a database write.')
doc.add_picture(str(ROOT/'docs/evidence/screenshots/leap-day-detail.png'),width=Inches(6.7))
p('Figure 1. Crop of the actual Playwright screenshot after reload: valid leap-day date and ₹125 transaction. Full capture and executable test are included.', 'Caption')

page();h('3. D2 / D3 — Request lifecycle failures')
sub('D2: the timeout stopped too early')
p('The transport test resolves fetch headers immediately with status 200, then makes JSON reading wait for an AbortSignal. With timeoutMs=50 and retries=0, advancing fake time to 60 ms should produce TIMEOUT. The original implementation remained pending. This failure is captured in red-client.log.')
p('Root cause: client/src/services/api.js called guard.done() after fetch resolved, before awaiting res.json(). Fetch resolves at headers, so a stalled body had no active deadline. The fix keeps JSON reading inside the guarded try block and preserves the abort reason. The timer/listener is released after the body completes or the failure is caught.')
table(['Boundary','Final expected result'],[
 ['Headers received; body stalls','TIMEOUT after the controlled 50 ms deadline'],
 ['204 DELETE response','Return null; do not call JSON reader'],
 ['Malformed successful JSON','HTTP_ERROR; no remaining timeout timer'],
 ['GET receives repeated 503','Three total attempts, then error'],
 ['Default POST network failure','One attempt; no automatic mutation retry'],
])
p('These are fake-timer/fetch unit tests, not real network latency measurements. The production timeout remains 12 seconds per attempt; retries mean the total GET operation can exceed 12 seconds.')
sub('D3: cancelled requests still entered a retry')
p('The second regression rejects the first fetch, aborts the caller after 10 ms during the 250 ms backoff, then advances timers. The old code invoked the fetch mock twice rather than once. Its sleep helper ignored cancellation, and the retry loop did not check for an already-aborted caller before invoking fetch again.')
p('The fix adds an abort-aware backoff Promise that clears its timer/listener and checks the caller signal at each loop entry. The final regression asserts exactly one fetch invocation and AbortError. Native fetch might reject an already-aborted signal before transmitting; therefore this evidence is described as an avoided invocation, not proof that two HTTP requests previously reached the server.')
sub('Why these tests matter')
p('The inherited tests exercised successful responses and broad network failures. They did not distinguish header completion from body completion or cancellation during a retry delay. Testing these transitions directly prevented false confidence from the otherwise passing baseline.')
code('Affected source: client/src/services/api.js\nRegression file: client/src/test/week5-api.test.js\nEvidence: red-client.log → final-client.log / client-results.json')

page();h('4. D4 — Credentials in request logs')
sub('Discovery')
p('Reviewing the first successful E2E output exposed SSE access logs containing /api/v1/events/stream?token=…. EventSource authentication passes the JWT in a query parameter, and the old Morgan dev/combined formatter logged the original URL. Saved evidence is redacted; no original JWT is included in the submission.')
sub('Fix and regression proof')
p('server/src/middleware/requestLogger.js now records method, path before the question mark, status, response time and request ID. It omits all query values and Referer rather than maintaining a fragile list of sensitive parameter names. server/src/app.js uses this formatter.')
code('Before (sanitized): GET /api/v1/events/stream?token=[REDACTED] 200 ...\nAfter: GET /api/v1/events/stream 200 ... requestId=...')
p('Four middleware tests first ran against the old Morgan format and failed because private-test-value appeared in output. They now pass for a normal token, repeated/mixed-case keys, percent-encoded keys and an arbitrary query value. Every case also supplies a sensitive Referer and requires route/status to remain visible.')
p('Trade-off: query-level diagnostics are reduced; request IDs retain traceability. The token still exists in the browser request and could appear in upstream proxy logs. Production proxy redaction and a reviewed cookie/short-lived ticket design remain follow-up work. The evidence-sanitizer script also removes JWTs from previously captured files before sharing.')
h('5. Environment debugging, not product defects')
table(['Failure observed','Action and outcome'],[
 ['Chromium could not load libnspr4.so','Installed Playwright Linux system dependencies; reran all 30 browser tests successfully.'],
 ['MongoDB: available disk below 524,288,000 bytes during index creation','Inspected the 993 MB /tmp volume; removed two confirmed stale 201 MB test DB directories with no active mongod; final 100-test server coverage run passed.'],
])
p('The storage failure produced 99 passing tests and one failed query-plan test before cleanup. It was not addressed by weakening the database safety threshold or removing the assertion. Both environment failure logs are retained separately from the defect evidence.')
p('Lesson: classify failures before changing application logic. A missing browser library or insufficient temporary disk space cannot establish an application regression.')

page();h('6. P1 — Measured query optimization')
sub('Profile-driven hypothesis')
p('The default transaction list sorts by date descending and _id descending to keep pages deterministic. Its old compound index included user and date only. MongoDB explain showed a blocking SORT and 10,000 examined documents for a 20-row page. The model index was extended to cover the entire sort; controller behaviour was not changed.')
code('// server/src/models/Transaction.js\n// Before: { user: 1, date: -1 }\ntransactionSchema.index({ user: 1, date: -1, _id: -1 })')
sub('Controlled method')
p('server/scripts/profile-transactions.mjs starts a dedicated ephemeral database, inserts 20,000 deterministic records for two owners (10,000 each) across 365 dates, then compares the reconstructed old index with the new index. No query hint is used. Each variant receives five warmups and 30 sequential measurements of a first-page, limit-20 native-driver query. Timings include query consumption but exclude HTTP, authentication and countDocuments. Percentiles use nearest-rank.')
profile=json.loads((ROOT/'docs/evidence/performance/query-profile.json').read_text())
b,a=profile['before'],profile['after']
table(['Metric','Old index','New index'],[
 ['Documents examined',b['totalDocsExamined'],a['totalDocsExamined']],
 ['Keys examined',b['totalKeysExamined'],a['totalKeysExamined']],
 ['Blocking SORT','Present','Absent'],
 ['Returned documents',20,20],
 ['Median query time',f"{b['medianMs']:.3f} ms",f"{a['medianMs']:.3f} ms"],
 ['p95 query time',f"{b['p95Ms']:.3f} ms",f"{a['p95Ms']:.3f} ms"],
 ['Same ordered IDs','Yes','Yes'],
])
p('All 60 raw timings, environment details and both explain plans are included. Two regression tests also check the application model index, no blocking SORT, 20 examined documents, disjoint equal-date pages and unchanged public serialization.')
sub('Limits and trade-offs')
p('This is a warm local database experiment, not a production-load or whole-dashboard speed claim. Old-first/new-second order may create cache bias; the structural explain improvement is stronger evidence than the timings. Wider indexes cost storage and write maintenance, neither benchmarked here. Deep skip pagination, alternate sorts, substring search and totals remain separate profiling targets.')
p('For existing persistent databases, create/verify the new index through a reviewed migration before optionally dropping the old one. Mongoose does not automatically remove the old index. No production database was modified. The front-end already uses lazy routes; its current entry bundle is 86.39 kB gzip and dashboard chunk 119.19 kB gzip. These are build sizes, not measured user latency, and no new bundle reduction is claimed.')

page();h('7. Test strategy, coverage and reproducibility')
table(['Layer','Tools / isolation','Why selected'],[
 ['Pure logic / transport','Vitest; controlled inputs, fake timers and fetch','Exercise hard-to-reproduce request transitions and boundaries.'],
 ['Components','Testing Library + jsdom','Assert behaviour, payloads, focus and accessible error linkage.'],
 ['API + database','Supertest + ephemeral MongoDB; cleanup after each test','Verify status contracts and actual persistence/isolation.'],
 ['Browser integration','Playwright Chromium; fresh users; one worker','Verify UI/API flows and reload persistence; inherited fault tests intercept requests intentionally.'],
])
p('The complete test inventory lists all 192 Vitest assertions and risk rationale. TEST-MATRIX.md maps the 40 additions to requirements and groups the inherited regression coverage. Browser tests are not described as entirely unmocked: selected inherited cases deliberately simulate outages and expired responses.')
table(['Source coverage','Lines','Branches','Interpretation'],[
 ['Server','89.72%','72.83%','100% validator lines; startup entry has 0% unit coverage.'],
 ['Client','79.77%','67.83%','DashboardPage 47.77%; ErrorBoundary 25% lines remain gaps.'],
])
p('Coverage includes source files with no unit execution. Browser execution is not merged into these V8 figures. Passing tests do not prove absence of defects; fallback rendering and dashboard edge states are worthwhile next additions.')
sub('Run locally in VS Code / PowerShell')
code('cd <path-to>\\week5-testing-optimization\nnpm.cmd run setup\nnpm.cmd test\nnpm.cmd run test:coverage\nnpm.cmd run lint\nnpm.cmd run build\ncd e2e\nnpx.cmd playwright install chromium\ncd ..\nnpm.cmd run test:e2e\nnpm.cmd run profile')
p('Use Node 22.12+ or 24 LTS; this evidence was generated with Node 20.20.2 on Linux. Windows instructions are supplied but not executed here. npm run setup installs locked dependencies with npm ci. E2E starts its own API/client, uses an ephemeral DB and refuses already-occupied ports. Stop existing app servers first. Linux may require Playwright --with-deps installation.')
p('To use the app: npm.cmd run dev, then http://localhost:5173/login; demo@spendwise.app / Demo1234. For persistent storage configure server/.env. Default demo data is temporary. The optional GitHub Actions template has not been executed on GitHub and is not claimed as CI evidence.')

page();h('8. Deliverables and evidence index')
table(['Requirement','Deliverable / verification'],[
 ['Compressed source + tests + config','SpendWise-Week5-Testing-Optimization.zip; client/, server/, e2e/, scripts/'],
 ['Unit and integration tests','TEST-MATRIX.md; TEST-INVENTORY.csv; final-server.log; final-client.log; final-e2e.log'],
 ['Error logs and debugging reports','DEBUGGING.md; red-*.log; environment-*.log; before/ snapshots; changes.patch'],
 ['Performance optimization evidence','PERFORMANCE.md; query-profile.json; before-explain.json; after-explain.json'],
 ['How to run / strategy / summary','Root README.md with exact commands, scope and troubleshooting'],
 ['Report upload','This .docx file; portal accepts Word documents'],
])
sub('Reading the evidence efficiently')
bullet('Start with baseline-tests.log to confirm the inherited server/client pass, then red-server.log, red-client.log, red-form.log and red-logging.log to see the exact missed assertions.')
bullet('Use final-server.log, final-client.log and final-e2e.log for the latest full runs. Machine-readable server-results.json and client-results.json expose each assertion result.')
bullet('Inspect query-profile.json for raw observations; the two explain plans show why work decreased. Do not infer total API response time from these query-only timings.')
bullet('The Week 5 screenshot is under docs/evidence/screenshots/. Older screenshots/video and WEEK4-README.md are retained only as historical Week 4 references.')
sub('Completion and remaining limits')
p('The final implementation passes 100 server, 92 client and 30 Chromium tests, lint and build checks. The work corrected calendar/range validation, response-body timeout lifetime, retry cancellation and credential-bearing request logs, and eliminated the measured blocking sort for the default list query. Each finding is tied to source files, a reproduction and a verification result.')
p('No Safari/Firefox/mobile certification, production deployment, concurrent load test or real-user latency measurement was performed. Upstream log redaction and persistent index rollout remain deployment responsibilities. There is no claim of 30–35 hours of actual work; that number is the assignment estimate, not a tracked time sheet.')
p('The main lesson is that a green happy-path suite needs targeted boundary tests and observable contracts. The report therefore retains failing evidence and limitations rather than replacing them with generic statements about improved quality.')
p('Submission URL: https://github.com/Badal00999/SmartSpend\nPush week5-testing-optimization/ before submitting the report link. GitHub upload and portal submission remain user actions; this report does not claim they have occurred.')

OUT.parent.mkdir(parents=True,exist_ok=True)
doc.save(OUT)
print(OUT)
